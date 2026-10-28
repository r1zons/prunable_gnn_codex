"""Author documents and plots exclusively from already archived results."""
from pathlib import Path
import sys, os, csv, json, re, statistics as st, hashlib
os.environ['MPLCONFIGDIR']='/private/tmp/sitito-mpl-config'
sys.path.insert(0,'/private/tmp/sitito-authoring-deps')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from latex2mathml.converter import convert as latex_mathml
from mathml2omml import convert as mathml_omml
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT=Path(__file__).resolve().parents[1]
def load(name):return list(csv.DictReader((OUT/name).open()))
data=load('verified_v3_results.csv'); rand=load('verified_random_results.csv')
ev=OUT/'evidence/research/experiments'
runs=list(csv.DictReader((ev/'local_structured_v3_sprint_20261027/results/runs.csv').open()))
random=list(csv.DictReader((ev/'local_structured_v2_q_random_pilot_20261027/random_trials.csv').open()))
timings=list(csv.DictReader((ev/'local_structured_v3_sprint_20261027/results/fixed_probe_timing.csv').open()))
refs=json.loads((OUT/'source/references.json').read_text())
plt.rcParams.update({'font.family':'DejaVu Serif','font.size':8,'axes.labelsize':8,'axes.titlesize':9,'legend.fontsize':8,'pdf.fonttype':42,'svg.fonttype':'none','axes.spines.top':False,'axes.spines.right':False})
styles={'q_reachable_v3':('Q','#3b6a9b','o'),'greedy_reachable_v3':('Greedy','#222222','x'),'random_reachable_v3':('Random действий','#a65d21','^')}
fig,axs=plt.subplots(2,2,figsize=(6.65,3.75),layout='constrained')
for ax,(m,t) in zip(axs.flat,[(m,t) for m in ['gcn','graphsage'] for t in [.3,.5]]):
    ax.axvspan(t*100-1,t*100+1,color='#e8eef3',zorder=0)
    for v,(label,color,marker) in styles.items():
        rr=[r for r in runs if r['platform']=='local' and r['model']==m and float(r['target'])==t and r['variant']==v]
        ax.scatter([float(r['achieved_sparsity'])*100 for r in rr],[float(r['validation_accuracy'])*100 for r in rr],label=label,c=color,marker=marker,s=35,linewidths=1.2,zorder=3 if marker=='x' else 2)
    ax.set(xlim=(t*100-1.15,t*100+1.15),ylim=(58,91),xlabel='Сокращение числа параметров, %',ylabel='Accuracy на валидации, %',title=f'{"GCN" if m=="gcn" else "GraphSAGE"}, цель {t*100:.0f} %')
    ax.grid(alpha=.18)
handles,labels=axs[0,0].get_legend_handles_labels()
fig.legend(handles,labels,loc='outside lower center',ncol=3,frameon=False)
for ext in ['png','pdf','svg']:fig.savefig(OUT/f'figures/fig1_results.{ext}',dpi=320)
plt.close(fig)
fig,axs=plt.subplots(2,2,figsize=(6.65,3.65),layout='constrained')
for ax,(m,t) in zip(axs.flat,[(m,t) for m in ['gcn','graphsage'] for t in [.3,.5]]):
    rr=[r for r in random if r['model']==m and float(r['target_sparsity'])==t]
    b=[json.loads(r['hidden_widths'])[1] for r in rr]
    ax.scatter(b,[float(r['validation_accuracy'])*100 for r in rr],s=25,color='#3b6a9b',alpha=.8,edgecolors='white',linewidths=.4)
    ax.set(xlim=(3,33),ylim=(0,100),xlabel='Ширина второго скрытого слоя b',ylabel='Accuracy на валидации, %',title=f'{"GCN" if m=="gcn" else "GraphSAGE"}, цель {t*100:.0f} %, a={23 if t==.3 else 16}')
    ax.grid(alpha=.18)
for ext in ['png','pdf','svg']:fig.savefig(OUT/f'figures/fig2_random.{ext}',dpi=320)
plt.close(fig)

def num(x,n=2):return f'{float(x):.{n}f}'.replace('.',',')
def ms(r,col):return num(float(r[col])*100)
def pm(r,c):
    a,b=float(r[c+'_mean'])*100,float(r[c+'_sd'])*100
    return num(a) if b<1e-10 else num(a)+' ± '+num(b)
variants=['r2_reference','q_stop_exact_v3','q_conservative_v3','q_random_ties_v3','q_target_stop_v3','q_reachable_v3','random_reachable_v3','greedy_reachable_v3']
names=['Q R2','Q точный STOP','Q консервативный','Q случайные равенства','Q пороговая остановка','Constrained Q','Constrained Random','Constrained Greedy']
hitrows=[]
for v,n in zip(variants,names):
    row=[n]
    for t in [.3,.5]:row.append(str(sum(int(r['hits']) for r in data if r['variant']==v and float(r['target'])==t))+'/6')
    row.append(str(sum(int(r['hits']) for r in data if r['variant']==v))+'/12');hitrows.append(row)
mainrows=[]
for m in ['gcn','graphsage']:
 for t in [.3,.5]:
  for v in ['q_reachable_v3','greedy_reachable_v3','random_reachable_v3']:
    r=next(r for r in data if r['model']==m and float(r['target'])==t and r['variant']==v)
    rr=[z for z in runs if z['platform']=='local' and z['model']==m and float(z['target'])==t and z['variant']==v]
    wp='\n'.join(dict.fromkeys(f"{z['width0']}×{z['width1']}; {int(z['parameters_after']):,}".replace(',',' ') for z in rr))
    mainrows.append([('GCN' if m=='gcn' else 'SAGE')+f'\n{t*100:.0f} %',styles[v][0].replace(' действий',''),wp,pm(r,'achieved_sparsity'),pm(r,'validation_accuracy'),pm(r,'validation_macro_f1'),num(r['search_seconds_mean'])])
randomrows=[[('GCN' if r['model']=='gcn' else 'GraphSAGE')+' '+str(round(float(r['target'])*100))+' %',str(r['n_masks']),num(float(r['mean'])*100)+' ± '+num(float(r['sd'])*100),num(float(r['minimum'])*100)+'–'+num(float(r['maximum'])*100)] for r in rand]
timingrows=[]
for platform in ['local','polus']:
 for m in ['gcn','graphsage']:
    rr=[r for r in timings if r['platform']==platform and r['model']==m]
    vals=[]
    for state in ['dense','fixed_pruned']:
        r=next(z for z in rr if z['model_state']==state)
        vals.append(num(float(r['mean_seconds'])*1000,3)+' ± '+num(float(r['std_seconds'])*1000,3))
    timingrows.append([('Локально' if platform=='local' else 'POLUS'),('GCN' if m=='gcn' else 'GraphSAGE'),*vals])
tables={'hits':(['Вариант','Цель 30 %','Цель 50 %','Всего'],hitrows,[7,3.3,3.3,3.4]),'main':(['Условие','Метод','Ширины; P','s, %','Acc., %','F1, %','Поиск, с'],mainrows,[1.7,1.55,2.9,2.7,2.7,2.7,1.75]),'random':(['Условие','Маски','Среднее ± SD, %','Минимум–максимум, %'],randomrows,[4,2,5,6]),'timing':(['Платформа','Модель','Dense, мс','После удаления, мс'],timingrows,[3,3,5.5,5.5])}
# Normalize widths to the 17 cm text area.
tables={k:(h,r,[v*17/sum(w) for v in w]) for k,(h,r,w) in tables.items()}

def reference(r,english=False):
    s=f"{r['authors']} {r['title']} // {r['venue']}. {r['year']}."
    if r.get('volume'):s+=(' Vol. ' if english else ' Т. ')+r['volume']+'.'
    if r.get('pages'):s+=(' P. ' if english else ' С. ')+r['pages']+'.'
    if r.get('doi'):
        label=' DOI'
        if r['doi_kind'].startswith('arXiv'):label+=' (arXiv)'
        s+=label+': '+r['doi']+'.'
    s+=' URL: '+r['url']+(' (accessed 27.10.2026).' if english else ' (дата обращения: 27.10.2026).')
    return s

doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor(0,0,0)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
pf=normal.paragraph_format;pf.line_spacing=1;pf.space_after=Pt(0);pf.space_before=Pt(0);pf.first_line_indent=Cm(1.25);pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.core_properties.title='Контроль достижимости параметрической компрессии при локальном структурном прунинге графовых нейронных сетей без дообучения'
doc.core_properties.author='[ЗАПОЛНИТЬ АВТОРАМ]';doc.core_properties.subject='Пилотное исследование Cora, validation-only, без новых научных экспериментов'

def paragraph(text='',bold=False,center=False,noindent=False,keep=False,before=0,after=0):
    p=doc.add_paragraph(style='Normal');p.paragraph_format.keep_with_next=keep;p.paragraph_format.widow_control=True
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    if noindent:p.paragraph_format.first_line_indent=Cm(0)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text);r.bold=bold
    if '[' in text and any(x in text for x in ['ЗАПОЛНИТЬ','ФАМИЛИЯ','АВТОР','FULL','SURNAME','НАЗВАНИЕ','УТОЧНИТЬ']):r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    return p

def table(key):
    headers,rows,widths=tables[key]
    tb=doc.add_table(rows=1, cols=len(headers));tb.autofit=False;tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,w in enumerate(widths):tb.columns[i].width=Cm(w)
    for i,h in enumerate(headers):tb.rows[0].cells[i].text=h
    for row in rows:
        cells=tb.add_row().cells
        for c,t in zip(cells,row):c.text=t
    for ri,row in enumerate(tb.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                z=OxmlElement('w:'+side);z.set(qn('w:w'),'55');z.set(qn('w:type'),'dxa');margin.append(z)
            cp.append(margin)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'B0B0B0');borders.append(z)
            cp.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cp.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=(key=='main' and ri>0 and ri%3!=0);p.paragraph_format.line_spacing=1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=ri==0
    paragraph('',noindent=True,after=2)

def escape(s):
    urls=[]
    def url_token(m):
        token=f'ZZURLTOKEN{len(urls)}ZZ';urls.append((token,r'\url{'+m.group(0)+'}'));return token
    s=re.sub(r'https?://[^\s]+',url_token,s)
    hashes=[]
    def hash_token(m):
        token=f'ZZHASHTOKEN{len(hashes)}ZZ';v=m.group(0)
        hashes.append((token,r'\texttt{'+r'\allowbreak '.join(v[i:i+8] for i in range(0,len(v),8))+'}'));return token
    s=re.sub(r'\b[0-9a-f]{40,64}\b',hash_token,s)
    chars={'\\':r'\textbackslash{}','{':r'\{','}':r'\}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    # Replace unicode sub/superscript glyphs unsupported in Times via standard inline math.
    replacements={'₀':'$_0$','ₜ':'$_t$','ℓ':'$\\ell$','Σⱼ₌₀ʰ⁻¹ γʲ':r'$\sum_{j=0}^{h-1}\gamma^j$','Φₜ':r'$\Phi_t$','Φₜ':'$\\Phi_t$','Φ':'$\\Phi$','→':r'$\to$','∈':r'$\in$','≤':r'$\leq$','≥':r'$\geq$','−':r'$-$','×':r'$\times$','±':r'$\pm$','⁻¹⁶':r'$^{-16}$','⁻¹⁵':r'$^{-15}$','⁻⁶':r'$^{-6}$','⁻⁵':r'$^{-5}$','α':r'$\alpha$','γ':r'$\gamma$','ε':r'$\varepsilon$','τ':r'$\tau$','λ':r'$\lambda$'}
    tokens=[]
    for a,b in sorted(replacements.items(),key=lambda kv:-len(kv[0])):
        if a in s:
            token=f'ZZREPLACEMENT{len(tokens)}ZZ';s=s.replace(a,token);tokens.append((token,b))
    s=''.join(chars.get(c,c) for c in s)
    for a,b in tokens:s=s.replace(a,b)
    for a,b in urls:s=s.replace(a,b)
    for a,b in hashes:s=s.replace(a,b)
    return s
def textext(s):
    parts=re.split(r'(\[\d+(?:,\s*\d+)*\])',s)
    return ''.join(r'\cite{'+','.join('r'+n.strip() for n in p[1:-1].split(','))+'}' if re.fullmatch(r'\[\d+(?:,\s*\d+)*\]',p) else escape(p) for p in parts)

tex=[r'''\documentclass[10pt,a4paper]{article}
\usepackage{fontspec}
\setmainfont{Times New Roman}
\usepackage[russian,english]{babel}
\usepackage[margin=2cm]{geometry}
\usepackage{amsmath,amssymb,graphicx,longtable,array,booktabs}
\usepackage[numbers,sort&compress]{natbib}
\usepackage[unicode,hidelinks]{hyperref}
\usepackage{url}
\Urlmuskip=0mu plus 2mu
\usepackage{titlesec}
\titleformat{\section}{\normalfont\bfseries\fontsize{10}{12}\selectfont}{}{0pt}{}
\titleformat{\subsection}{\normalfont\bfseries\fontsize{10}{12}\selectfont}{}{0pt}{}
\titlespacing*{\section}{0pt}{7pt}{3pt}
\titlespacing*{\subsection}{0pt}{5pt}{2pt}
\setlength{\parindent}{1.25cm}
\setlength{\parskip}{0pt}
\setlength{\emergencystretch}{3em}
\pagestyle{empty}
\begin{document}
\selectlanguage{russian}
''']
md=[];eq=0
for raw in (OUT/'source/manuscript.md').read_text().splitlines():
    s=raw.strip()
    if not s:continue
    if s.startswith('@'):
        parts=s.split('|',2);kind=parts[0]
        if kind=='@eq':
            eq+=1;formula=s.split('|',1)[1].strip()
            p=paragraph('',noindent=True,center=True,before=3,after=3)
            omml=mathml_omml(latex_mathml(formula))
            p._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'+re.sub(r'^<m:oMath[^>]*>|</m:oMath>$','',omml)+'</m:oMath>') if not omml.startswith('<m:oMath') else parse_xml(omml.replace('<m:oMath>','<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">',1)))
            p.add_run('   ('+str(eq)+')')
            tex.append('\\begin{equation}\n'+formula+'\n\\end{equation}\n');md.append('$$'+formula+'$$');continue
        if kind=='@table':
            key,caption=parts[1:];paragraph(caption,noindent=True,bold=False,keep=True,before=5,after=3);table(key)
            headers,rows,widths=tables[key]
            cols=''.join('p{'+f'{w/17*.94:.3f}'+r'\linewidth}' for w in widths)
            tex.append('\\par\\medskip\\noindent '+escape(caption)+'\\par\n\\begingroup\\setlength{\\tabcolsep}{2pt}\\renewcommand{\\arraystretch}{1.12}\n\\begin{longtable}{'+cols+'}\n\\toprule\n'+' & '.join(escape(h) for h in headers)+r'\\ \midrule\endhead'+'\n')
            for row in rows:tex.append(' & '.join(escape(x.replace('\n',' / ')) for x in row)+r'\\'+'\n')
            tex.append(r'\bottomrule\end{longtable}\endgroup'+'\n')
            md.extend([caption,'| '+' | '.join(headers)+' |','|'+'---|'*len(headers)]+['| '+' | '.join(z.replace('\n',' / ') for z in row)+' |' for row in rows]);continue
        if kind=='@figure':
            file,caption=parts[1:];p=paragraph('',center=True,noindent=True,keep=True,before=5)
            p.add_run().add_picture(str(OUT/'figures'/file),width=Cm(16.5))
            paragraph(caption,noindent=True,center=True,before=2,after=5)
            tex.append('\\begin{figure}[htbp]\\centering\n\\includegraphics[width=\\linewidth]{../figures/'+file.replace('.png','.pdf')+'}\n\\par\\smallskip '+escape(caption)+'\n\\end{figure}\n');md.extend(['![График](figures/'+file+')',caption]);continue
        if kind=='@bibliography':
            paragraph('Список литературы',bold=True,noindent=True,keep=True,before=7,after=3)
            for r in refs:paragraph(str(r['n'])+'. '+reference(r),noindent=True,after=2)
            paragraph('References',bold=True,noindent=True,keep=True,before=7,after=3)
            for r in refs:paragraph(str(r['n'])+'. '+reference(r,True),noindent=True,after=2)
            tex.append(r'\renewcommand{\refname}{Список литературы}\bibliographystyle{unsrtnat}\bibliography{references}'+'\n')
            tex.append(r'\section*{References}'+'\n')
            for r in refs:tex.append('\\noindent '+str(r['n'])+'. '+escape(reference(r,True))+'\\par\n')
            md.extend(['## Список литературы']+[str(r['n'])+'. '+reference(r) for r in refs]+['## References']+[str(r['n'])+'. '+reference(r,True) for r in refs]);continue
        text=parts[1]
        if kind in ['@title','@title_en']:paragraph(text,bold=True,center=True,noindent=True,before=5,after=5);tex.append('\\begin{center}\\bfseries '+escape(text)+'\\end{center}\n')
        elif kind.startswith('@abstract'):
            label='Аннотация. ' if kind.endswith('_ru') else 'Abstract. '
            paragraph(label+text,noindent=True,after=4);tex.append('\\noindent '+escape(label+text)+'\\par\n')
        elif kind.startswith('@keywords'):
            label='Ключевые слова: ' if kind.endswith('_ru') else 'Keywords: '
            paragraph(label+text,noindent=True,after=5);tex.append('\\noindent '+escape(label+text)+'\\par\n')
        else:
            paragraph(text,noindent=True,center=kind in ['@author','@affiliation'],after=2)
            tex.append('\\noindent '+escape(text)+'\\par\n')
        md.append(text);continue
    if s.startswith('#'):
        title=s.lstrip('# ');hp=paragraph(title,bold=True,noindent=True,keep=True,before=7,after=3)
        if title=='Введение':hp.paragraph_format.page_break_before=True;tex.append('\\clearpage\n')
        tex.append(('\\subsection*{' if s.startswith('###') else '\\section*{')+escape(title)+'}\n');md.append(s);continue
    pp=paragraph(s)
    if s.startswith('Графовые нейронные сети используются'):
        run=pp.add_run();ref=OxmlElement('w:footnoteReference');ref.set(qn('w:id'),'1');run._r.append(ref)
    tex.append(textext(s)+('\\footnote{GNN: graph neural network; GCN: graph convolutional network; GraphSAGE: graph sample and aggregate; CPU: central processing unit; SD: sample standard deviation; DP: dynamic programming; PyG: PyTorch Geometric.}' if s.startswith('Графовые нейронные сети используются') else '')+'\n\n');md.append(s)
tex.append('\\end{document}\n')
footnote_text='GNN — graph neural network, графовая нейронная сеть; GCN — graph convolutional network; GraphSAGE — graph sample and aggregate; CPU — центральный процессор; SD — выборочное стандартное отклонение; DP — динамическое программирование; PyG — PyTorch Geometric. Acc. — accuracy, доля правильных ответов; macro-F1 — невзвешенное среднее F1 по семи классам.'
footnotes=OxmlElement('w:footnotes')
for fid,tag in [('-1','separator'),('0','continuationSeparator')]:
    f=OxmlElement('w:footnote');f.set(qn('w:type'),tag);f.set(qn('w:id'),fid);p=OxmlElement('w:p');r=OxmlElement('w:r');r.append(OxmlElement('w:'+tag));p.append(r);f.append(p);footnotes.append(f)
f=OxmlElement('w:footnote');f.set(qn('w:id'),'1');p=OxmlElement('w:p');pp=OxmlElement('w:pPr');pp.append(OxmlElement('w:keepLines'));p.append(pp);r=OxmlElement('w:r');r.append(OxmlElement('w:footnoteRef'));p.append(r)
r=OxmlElement('w:r');rp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');rp.append(fonts);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');rp.append(sz);r.append(rp);tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=' '+footnote_text;r.append(tt);p.append(r);f.append(p);footnotes.append(f)
from lxml import etree
fp=Part(PackURI('/word/footnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',etree.tostring(footnotes),doc.part.package)
doc.part.relate_to(fp,RT.FOOTNOTES)
doc.save(OUT/'manuscript_sitito.docx')
(OUT/'source/main.tex').write_text(''.join(tex))
(OUT/'manuscript_readable.md').write_text('\n\n'.join(md))
bib=[]
for r in refs:
    kind='misc' if r['n']==16 else ('article' if r['n']==13 else 'inproceedings')
    venuefield='journal' if kind=='article' else ('howpublished' if kind=='misc' else 'booktitle')
    fields={'author':r['bibauthors'],'title':'{'+r['title'].replace('_',r'\_')+'}','year':str(r['year']),venuefield:r['venue'].replace('_',r'\_'),'url':r['url']}
    for key in ['volume','pages','doi']:
        if r.get(key):fields[key]=r[key].replace('–','--')
    if r.get('doi','').startswith('10.48550'):fields['note']='DOI identifies the arXiv electronic version'
    bib.append('@'+kind+'{r'+str(r['n'])+',\n'+',\n'.join('  '+k+' = {'+v+'}' for k,v in fields.items())+'\n}\n')
(OUT/'source/references.bib').write_text('\n'.join(bib))
(OUT/'source/tables.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2))
print('Created DOCX, TeX, BibTeX, two figures as PNG/PDF/SVG; equations:',eq)
